"""Lista ofert ze statusami: scalanie, zakres przebiegu, wspólny status, arkusz."""

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import companies_main
import portals_main
import public_main
import report


def row(company, title, link, status="Nowa", active="Tak", first="2026-10-01"):
    return {"Firma": company, "Stanowisko": title, "Link": link, "Status": status,
            "W ostatnim wyszukiwaniu": active, "Pierwsze wykrycie": first}


def current(company, title, link):
    return {"Firma": company, "Stanowisko": title, "Link": link}


def test_merge_keeps_status_and_marks_new_and_gone():
    previous = [row("A", "QA Tester", "https://a/1", status="CV wysłane"), row("C", "Tester", "https://c/1")]
    rows, new = report._merge(previous, [current("A", "QA Tester", "https://a/1"), current("D", "Tester", "https://d/1")],
                              "2026-10-02")
    by_company = {item["Firma"]: item for item in rows}
    assert by_company["A"]["Status"] == "CV wysłane"
    assert by_company["A"]["Pierwsze wykrycie"] == "2026-10-01"
    assert [item["Firma"] for item in new] == ["D"]
    assert by_company["C"]["W ostatnim wyszukiwaniu"] == "Nie"


def test_merge_matches_by_company_and_title_when_link_changes():
    previous = [row("A", "QA Tester", "https://a/1", status="Obejrzana")]
    rows, new = report._merge(previous, [current("A", "QA Tester", "https://a/1?ref=x")], "2026-10-02")
    assert rows[0]["Status"] == "Obejrzana" and new == []


def test_merge_scope_keeps_unchecked_companies():
    # Regresja: --company Sii oznaczało oferty wszystkich innych firm jako zniknięte.
    previous = [row("Sii", "QA", "https://s/1"), row("Xebia", "QA 2", "https://x/1")]
    rows, _ = report._merge(previous, [], "2026-10-02", scope={"Sii"})
    assert {item["Firma"]: item["W ostatnim wyszukiwaniu"] for item in rows} == {"Sii": "Nie", "Xebia": "Tak"}


def test_company_core_ignores_legal_forms():
    assert report.company_core("Sii Polska Sp. z o.o.") == report.company_core("Sii") == "sii"
    assert report.company_core("Asseco Poland S.A.") == "asseco"


def test_shared_status_from_other_list_wins_when_more_advanced():
    previous = [row("Sii", "Manual Tester (f/m/x)", "https://sii.pl/1")]
    shared = {("sii", "manual tester"): "CV wysłane"}
    rows, _ = report._merge(previous, [], "2026-10-02", shared=shared)
    assert rows[0]["Status"] == "CV wysłane"


def test_shared_status_does_not_downgrade():
    previous = [row("Sii", "Manual Tester", "https://sii.pl/1", status="CV wysłane")]
    rows, _ = report._merge(previous, [], "2026-10-02", shared={"https://sii.pl/1": "Obejrzana"})
    assert rows[0]["Status"] == "CV wysłane"


def test_report_tabs_match_scripts():
    expected = {
        portals_main.SHEET_TAB: f"{portals_main.REPORTS_DIR}/{portals_main.BASENAME}.xlsx",
        companies_main.SHEET_TAB: f"{companies_main.REPORTS_DIR}/{companies_main.BASENAME}.xlsx",
        public_main.SHEET_TAB: f"{public_main.REPORTS_DIR}/{public_main.BASENAME}.xlsx",
    }
    assert {tab: path.as_posix() for tab, path in report.REPORT_TABS.items()} == expected


def test_excel_title_is_link_and_link_column_hidden(tmp_path, make_offer, monkeypatch):
    monkeypatch.setattr(report, "REPORT_TABS", {})
    result = report.export([make_offer(url="https://example.com/1", title="Manual Tester")], tmp_path, "t")
    sheet = load_workbook(result.excel_path).active
    header = [cell.value for cell in sheet[1]]
    assert "Notatka" not in header and "Manual / automat" in header
    title_cell = sheet.cell(row=2, column=header.index("Stanowisko") + 1)
    assert title_cell.hyperlink.target == "https://example.com/1"
    assert sheet.column_dimensions[get_column_letter(header.index("Link") + 1)].hidden


class FakeStore:
    class settings:  # noqa: N801 - atrapa atrybutu SheetStore.settings
        url = "https://docs.google.com/spreadsheets/d/FAKE"

    def __init__(self):
        self.tabs, self.fail_read, self.writes = {}, False, 0

    def read_rows(self, tab):
        if self.fail_read:
            raise ConnectionError("brak sieci")
        return [dict(item) for item in self.tabs.get(tab, [])]

    def write_rows(self, tab, columns, rows, *args, **kwargs):
        self.tabs[tab] = [dict(item) for item in rows]
        self.writes += 1


def test_sheet_status_survives_and_read_failure_blocks_write(tmp_path, make_offer, monkeypatch):
    store = FakeStore()
    monkeypatch.setattr(report, "_STORE_CACHE", {"store": store})
    monkeypatch.setattr(report, "REPORT_TABS", {})
    offers = [make_offer(company="A", title="QA Tester", url="https://a/1")]
    report.export(offers, tmp_path, "t", sheet_tab="Firmy")
    store.tabs["Firmy"][0]["Status"] = "CV wysłane"  # zmiana "z telefonu"
    report.export(offers, tmp_path, "t", sheet_tab="Firmy")
    assert store.tabs["Firmy"][0]["Status"] == "CV wysłane"

    store.fail_read, writes = True, store.writes
    result = report.export(offers, tmp_path, "t", sheet_tab="Firmy")
    assert store.writes == writes  # bez odczytu nie nadpisujemy statusów w arkuszu
    assert result.warnings and result.sheet_url == ""
