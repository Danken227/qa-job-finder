"""Eksport ofert do Excela (trwała lista ze statusami) i HTML.

Plik Excel jest jednocześnie "trackerem": przy każdym uruchomieniu program
wczytuje poprzednią wersję i przenosi to, co wpisałeś ręcznie - kolumny
**Status** i **Notatka** oraz datę pierwszego wykrycia. Oferty, których nie
ma w bieżącym wyszukiwaniu, nie znikają - dostają "Nie" w kolumnie
"W ostatnim wyszukiwaniu" (zamknięte, zmienione albo odfiltrowane).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from html import escape
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from collectors.base import JobOffer
from collectors.parsing import normalize_title

STATUS_NEW = "Nowa"
STATUSES = (STATUS_NEW, "Obejrzana", "CV wysłane", "Rozmowa", "Odrzucona", "Nie interesuje mnie")
# Statusy "załatwione" - wyszarzone w Excelu, zwinięte w HTML.
DONE_STATUSES = ("CV wysłane", "Rozmowa", "Odrzucona", "Nie interesuje mnie")
STATUS_ORDER = {status: index for index, status in enumerate(STATUSES)}

USER_COLUMNS = ("Status", "Notatka")
COLUMNS = (
    "Status",
    "Notatka",
    "Priorytet",
    "Firma",
    "Stanowisko",
    "Lokalizacja",
    "Model pracy",
    "Umowa",
    "Wynagrodzenie",
    "Ocena",
    "Dlaczego pasuje",
    "Ocena wynagrodzenia",
    "Źródło",
    "Pierwsze wykrycie",
    "Ostatnio widziana",
    "W ostatnim wyszukiwaniu",
    "Link",
)
WIDTHS = {
    "Status": 16, "Notatka": 30, "Priorytet": 18, "Firma": 24, "Stanowisko": 40, "Lokalizacja": 22,
    "Model pracy": 13, "Umowa": 18, "Wynagrodzenie": 24, "Ocena": 8, "Dlaczego pasuje": 36,
    "Ocena wynagrodzenia": 30, "Źródło": 18, "Pierwsze wykrycie": 13, "Ostatnio widziana": 13,
    "W ostatnim wyszukiwaniu": 12, "Link": 60,
}


@dataclass(slots=True)
class ExportResult:
    excel_path: Path
    html_path: Path
    new_rows: list[dict[str, str]] = field(default_factory=list)
    active_rows: int = 0
    warning: str = ""


def export(
    offers: list[JobOffer],
    reports_dir: Path | str = "reports",
    basename: str = "report",
    title: str = "Raport ofert QA",
) -> ExportResult:
    """Scala bieżące oferty z poprzednim Excelem i zapisuje Excel + HTML."""

    output_dir = Path(reports_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    excel_path = output_dir / f"{basename}.xlsx"
    html_path = output_dir / f"{basename}.html"
    today = date.today().isoformat()

    previous = _read_previous(excel_path)
    rows, new_rows = _merge(previous, [_to_row(offer) for offer in offers], today)

    warning = ""
    try:
        _write_excel(rows, excel_path)
    except PermissionError:
        # Plik jest otwarty w Excelu - nie tracimy wyników, zapisujemy kopię.
        stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
        fallback = output_dir / f"{basename} (kopia {stamp}).xlsx"
        _write_excel(rows, fallback)
        warning = (
            f"{excel_path} jest otwarty w Excelu - zapisano kopię {fallback.name}. "
            "Zamknij plik przed kolejnym uruchomieniem; statusy wpisane w kopii nie zostaną przeniesione."
        )
        excel_path = fallback
    _write_html(rows, html_path, title, today)
    active = sum(1 for row in rows if row["W ostatnim wyszukiwaniu"] == "Tak")
    return ExportResult(excel_path, html_path, new_rows, active, warning)


def _to_row(offer: JobOffer) -> dict[str, str]:
    priority = "🟢 Aplikuj" if offer.match_score >= 12 else "🟡 Warto rozważyć"
    reasons = ", ".join(offer.match_reasons) or "Dopasowane stanowisko QA"
    score_on_ten = max(1, min(10, round(offer.match_score / 2)))
    return {
        "Priorytet": priority,
        "Firma": offer.company,
        "Stanowisko": offer.title,
        "Lokalizacja": offer.location,
        "Model pracy": offer.work_mode,
        "Umowa": ", ".join(offer.contract_types) or "Nie podano",
        "Wynagrodzenie": offer.salary or "Nie podano",
        "Ocena": f"{score_on_ten}/10",
        "Dlaczego pasuje": reasons,
        "Ocena wynagrodzenia": offer.salary_assessment,
        "Źródło": offer.source,
        "Link": offer.url,
    }


# --- Scalanie z poprzednim raportem ------------------------------------------

def _key(row: dict[str, str]) -> tuple[str, str]:
    return (str(row.get("Firma", "")).casefold().strip(), normalize_title(str(row.get("Stanowisko", ""))))


def _read_previous(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    try:
        sheet = load_workbook(path, read_only=True).active
    except Exception:  # noqa: BLE001 - uszkodzony plik nie może zablokować raportu
        return []
    rows = sheet.iter_rows(values_only=True)
    header = [str(value or "") for value in next(rows, ())]
    if "Link" not in header:
        return []
    previous = []
    for values in rows:
        row = {column: ("" if value is None else str(value)) for column, value in zip(header, values)}
        if row.get("Link"):
            previous.append(row)
    return previous


def _merge(
    previous: list[dict[str, str]],
    current: list[dict[str, str]],
    today: str,
) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    """Zwraca (wszystkie wiersze, oferty nowe w tym przebiegu)."""

    by_url = {row["Link"]: row for row in previous}
    by_key = {_key(row): row for row in previous}
    matched: set[int] = set()
    rows: list[dict[str, str]] = []
    new_rows: list[dict[str, str]] = []

    for row in current:
        # Link bywa zmieniany przez portal; firma + tytuł to klucz zapasowy.
        old = by_url.get(row["Link"]) or by_key.get(_key(row))
        if old is not None and id(old) not in matched:
            matched.add(id(old))
            row["Status"] = old.get("Status") or STATUS_NEW
            row["Notatka"] = old.get("Notatka", "")
            row["Pierwsze wykrycie"] = old.get("Pierwsze wykrycie") or today
        else:
            row["Status"] = STATUS_NEW
            row["Notatka"] = ""
            row["Pierwsze wykrycie"] = today
            new_rows.append(row)
        row["Ostatnio widziana"] = today
        row["W ostatnim wyszukiwaniu"] = "Tak"
        rows.append(row)

    for old in previous:
        if id(old) in matched:
            continue
        old = {column: old.get(column, "") for column in COLUMNS}
        old["W ostatnim wyszukiwaniu"] = "Nie"
        rows.append(old)

    # Najpierw najnowsze wykrycia, potem (sortowanie stabilne) aktywne i
    # niezałatwione na górze.
    rows.sort(key=lambda row: row.get("Pierwsze wykrycie", ""), reverse=True)
    rows.sort(
        key=lambda row: (
            row["W ostatnim wyszukiwaniu"] != "Tak",
            row.get("Status") in DONE_STATUSES,
            STATUS_ORDER.get(row.get("Status", ""), len(STATUSES)),
        )
    )
    return rows, new_rows


# --- Excel -------------------------------------------------------------------

def _write_excel(rows: list[dict[str, str]], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Oferty"
    sheet.append(list(COLUMNS))
    for row in rows:
        sheet.append([row.get(column, "") for column in COLUMNS])

    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    user_fill = PatternFill("solid", fgColor="C55A11")
    for index, cell in enumerate(sheet[1], start=1):
        cell.fill = user_fill if COLUMNS[index - 1] in USER_COLUMNS else header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, column in enumerate(COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = WIDTHS[column]

    link_column = COLUMNS.index("Link") + 1
    for row_index in range(2, sheet.max_row + 1):
        link = sheet.cell(row=row_index, column=link_column)
        link.hyperlink = link.value
        link.style = "Hyperlink"
        for cell in sheet[row_index]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # Lista rozwijana statusów - również dla kilkuset pustych wierszy poniżej.
    status_letter = get_column_letter(COLUMNS.index("Status") + 1)
    last_row = sheet.max_row + 500
    validation = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=True)
    validation.error = "Wybierz status z listy."
    validation.prompt = "Zmień status - zostanie zachowany przy kolejnym uruchomieniu."
    sheet.add_data_validation(validation)
    validation.add(f"{status_letter}2:{status_letter}{last_row}")

    # Formatowanie warunkowe działa na bieżąco - po zmianie statusu wiersz od razu szarzeje.
    full_range = f"A2:{get_column_letter(len(COLUMNS))}{last_row}"
    active_letter = get_column_letter(COLUMNS.index("W ostatnim wyszukiwaniu") + 1)
    done_formula = " , ".join(f'${status_letter}2="{status}"' for status in DONE_STATUSES)
    sheet.conditional_formatting.add(
        full_range,
        FormulaRule(formula=[f'${active_letter}2="Nie"'], font=Font(color="A6A6A6", strike=True), stopIfTrue=True),
    )
    sheet.conditional_formatting.add(
        full_range,
        FormulaRule(formula=[f"OR({done_formula})"], font=Font(color="808080"),
                    fill=PatternFill("solid", fgColor="EDEDED"), stopIfTrue=True),
    )
    sheet.conditional_formatting.add(
        full_range,
        FormulaRule(formula=[f'${status_letter}2="{STATUS_NEW}"'], fill=PatternFill("solid", fgColor="E2EFDA")),
    )
    workbook.save(path)


# --- HTML --------------------------------------------------------------------

HTML_COLUMNS = ("Status", "Priorytet", "Firma", "Stanowisko", "Lokalizacja", "Model pracy", "Umowa",
                "Wynagrodzenie", "Ocena", "Dlaczego pasuje", "Pierwsze wykrycie", "Link")


def _write_html(rows: list[dict[str, str]], html_path: Path, title: str, today: str) -> None:
    active = [row for row in rows if row["W ostatnim wyszukiwaniu"] == "Tak" and row["Status"] not in DONE_STATUSES]
    done = [row for row in rows if row["W ostatnim wyszukiwaniu"] == "Tak" and row["Status"] in DONE_STATUSES]
    gone = [row for row in rows if row["W ostatnim wyszukiwaniu"] != "Tak"]
    new_count = sum(1 for row in active if row["Pierwsze wykrycie"] == today)

    sections = [_html_table(active, today, "Brak nowych ani obejrzanych ofert spełniających kryteria.")]
    if done:
        sections.append(f"<details><summary>Załatwione ({len(done)})</summary>{_html_table(done, today)}</details>")
    if gone:
        sections.append(
            f"<details><summary>Nie ma ich w ostatnim wyszukiwaniu ({len(gone)})</summary>{_html_table(gone, today)}</details>"
        )

    html_path.write_text(
        f"""<!doctype html>
<html lang="pl">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{escape(title)}</title>
  <style>
    body {{ font-family: Arial, sans-serif; margin: 2rem; color: #172033; }}
    h1 {{ margin-bottom: .25rem; }}
    table {{ border-collapse: collapse; width: 100%; font-size: .9rem; margin-bottom: 1.5rem; }}
    th {{ background: #1f4e78; color: white; text-align: left; position: sticky; top: 0; }}
    th, td {{ border: 1px solid #d7dee9; padding: .6rem; vertical-align: top; }}
    tr:nth-child(even) {{ background: #f7f9fc; }}
    tr.new td:first-child {{ background: #e2efda; font-weight: 600; }}
    details {{ margin-top: 1rem; }} summary {{ cursor: pointer; font-weight: 600; margin-bottom: .5rem; }}
    details td {{ color: #707070; }}
    a {{ color: #0b63ce; font-weight: 600; }}
  </style>
</head>
<body>
  <h1>{escape(title)}</h1>
  <p>Do przejrzenia: {len(active)} (nowe dziś: {new_count}). Każdy link został sprawdzony przed dodaniem do raportu.
  Statusy zmieniasz w pliku Excel obok - ten widok odświeży się przy kolejnym uruchomieniu.</p>
  {"".join(sections)}
</body>
</html>
""",
        encoding="utf-8",
    )


def _html_table(rows: list[dict[str, str]], today: str, empty: str = "") -> str:
    header = "".join(f"<th>{escape(column)}</th>" for column in HTML_COLUMNS)
    if not rows:
        return f'<table><thead><tr>{header}</tr></thead><tbody><tr><td colspan="{len(HTML_COLUMNS)}">{escape(empty)}</td></tr></tbody></table>'
    body = []
    for row in rows:
        cells = []
        for column in HTML_COLUMNS:
            value = str(row.get(column, ""))
            if column == "Link":
                cells.append(f'<td><a href="{escape(value, quote=True)}" target="_blank" rel="noopener">Otwórz ofertę</a></td>')
            else:
                cells.append(f"<td>{escape(value)}</td>")
        css = ' class="new"' if row.get("Pierwsze wykrycie") == today and row.get("Status") == STATUS_NEW else ""
        body.append(f"<tr{css}>{''.join(cells)}</tr>")
    return f"<table><thead><tr>{header}</tr></thead><tbody>{''.join(body)}</tbody></table>"
